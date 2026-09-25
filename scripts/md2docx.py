"""Markdown to Word, for the two documents the legal department reads.

Supports headings, paragraphs, bullets, numbered lists, pipe tables, block
quotes as callouts, horizontal rules as page breaks, and inline bold, italic
and code. Nothing else, because anything else would be a feature of the
converter rather than of the document.
"""

from __future__ import annotations

import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

INDIGO = RGBColor(0x3E, 0x40, 0x95)
DEEP = RGBColor(0x01, 0x4D, 0xBD)
INK = RGBColor(0x16, 0x17, 0x2B)
MUTED = RGBColor(0x6E, 0x70, 0x89)
BODY_FONT = "Calibri"
MONO_FONT = "Consolas"


def shade(cell, hexcolour: str) -> None:
    element = OxmlElement("w:shd")
    element.set(qn("w:val"), "clear")
    element.set(qn("w:fill"), hexcolour)
    cell._tc.get_or_add_tcPr().append(element)


def repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def borders(table, colour: str = "C8CAD8") -> None:
    properties = table._tbl.tblPr
    element = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        line = OxmlElement(f"w:{edge}")
        line.set(qn("w:val"), "single")
        line.set(qn("w:sz"), "4")
        line.set(qn("w:color"), colour)
        element.append(line)
    properties.append(element)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+?\*)")


def write_runs(paragraph, text: str, size: float = 10.0, colour=INK) -> None:
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            run = paragraph.add_run(piece[2:-2])
            run.bold = True
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = MONO_FONT
            run.font.size = Pt(size - 1)
            run.font.color.rgb = DEEP
            continue
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            run = paragraph.add_run(piece[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(piece)
        run.font.name = BODY_FONT
        run.font.size = Pt(size)
        run.font.color.rgb = colour


def base_document() -> Document:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.06
    for name, size, colour, before in (
        ("Title", 24, INDIGO, 0),
        ("Heading 1", 15, INDIGO, 13),
        ("Heading 2", 12, DEEP, 9),
        ("Heading 3", 10.5, INK, 7),
        ("Heading 4", 10, MUTED, 6),
    ):
        style = document.styles[name]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.color.rgb = colour
        style.font.bold = name != "Title"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.keep_with_next = True
    return document


def add_table(document, rows: list[list[str]]) -> None:
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = True
    borders(table)
    for index, row in enumerate(rows):
        for column, text in enumerate(row):
            cell = table.cell(index, column)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(1)
            paragraph.paragraph_format.space_before = Pt(1)
            write_runs(paragraph, text, size=8.8)
            if index == 0:
                shade(cell, "EEEFF4")
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = DEEP
    repeat_header(table.rows[0])
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_callout(document, lines: list[str]) -> None:
    table = document.add_table(rows=1, cols=1)
    borders(table, "9CD0FB")
    cell = table.cell(0, 0)
    shade(cell, "E6F3FE")
    cell.text = ""
    for index, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(3)
        write_runs(paragraph, line, size=9.5)
    document.add_paragraph().paragraph_format.space_after = Pt(4)


def convert(source: str, output: str) -> None:
    document = base_document()
    lines = source.split("\n")
    index = 0
    first_heading = True
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if not stripped:
            index += 1
            continue

        if stripped == "---":
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            index += 1
            continue

        if stripped.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [c.strip() for c in lines[index].strip().strip("|").split("|")]
                if not all(set(c) <= set("-: ") and c for c in cells):
                    rows.append(cells)
                index += 1
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            add_table(document, rows)
            continue

        if stripped.startswith(">"):
            block = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                block.append(lines[index].strip().lstrip(">").strip())
                index += 1
            add_callout(document, block)
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            if level == 1 and first_heading:
                paragraph = document.add_paragraph(style="Title")
                write_runs(paragraph, text, size=24, colour=INDIGO)
                for run in paragraph.runs:
                    run.font.bold = True
                first_heading = False
            else:
                style = {2: "Heading 1", 3: "Heading 2", 4: "Heading 3"}.get(level, "Heading 4")
                document.add_paragraph(text, style=style)
            index += 1
            continue

        if re.match(r"^[-*] ", stripped):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(2)
            write_runs(paragraph, stripped[2:])
            index += 1
            continue

        if re.match(r"^\d+\. ", stripped):
            paragraph = document.add_paragraph(style="List Number")
            paragraph.paragraph_format.space_after = Pt(3)
            write_runs(paragraph, re.sub(r"^\d+\. ", "", stripped))
            index += 1
            continue

        paragraph = document.add_paragraph()
        write_runs(paragraph, stripped)
        index += 1

    section = document.sections[0]
    for attribute in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, attribute, Inches(0.8))
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.text = ""
    run = footer.add_run("Data Science Nigeria and EqualyzAI, internal and confidential")
    run.font.size = Pt(8)
    run.font.name = BODY_FONT
    run.font.color.rgb = MUTED
    document.save(output)


if __name__ == "__main__":
    convert(open(sys.argv[1]).read(), sys.argv[2])
    print(f"wrote {sys.argv[2]}")
